#!/usr/bin/env python3
"""Every warning the tool can raise actually fires (and the file is still written).

Run: python3 tests/test_warnings.py   (needs requirements-dev.txt)
"""

import io
import tempfile
import zipfile
from pathlib import Path

from helpers import Checker, run

c = Checker("warnings")


def add_entries(src: Path, dst: Path, entries: dict[str, bytes]) -> Path:
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            zout.writestr(info, zin.read(info))
        for name, data in entries.items():
            zout.writestr(name, data)
    return dst


def expect_warning(label, result, out_dir, fragment, wrote=True):
    c.check(result.returncode == 2, f"{label}: exit status 2", f"rc={result.returncode} {result.stderr}")
    c.check(fragment in result.stderr, f"{label}: warning mentions {fragment!r}", result.stderr)
    if wrote:
        c.check(any(out_dir.iterdir()), f"{label}: output still written")


def png_bytes() -> bytes:
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (60, 30), "white").save(buf, "PNG")
    return buf.getvalue()


with tempfile.TemporaryDirectory() as tmp:
    td = Path(tmp)
    terms = ["--scrub", "Acme Corp"]

    import docx
    d = docx.Document()
    d.add_paragraph("Hello Acme Corp")
    base_docx = td / "base.docx"
    d.save(base_docx)

    def office_case(label, entries, fragment):
        src = add_entries(base_docx, td / f"{label}.docx".replace(" ", "_"), entries)
        out = td / f"out_{abs(hash(label))}"
        expect_warning(label, run(src, *terms, "--output-dir", out), out, fragment)

    office_case("embedded OLE object", {"word/embeddings/oleObject1.bin": b"\0\1\2"}, "embedded object")
    office_case("unscrubbable EMF image", {"word/media/image9.emf": b"\x01\0\0\0"}, ".emf image not scrubbed")
    office_case("undecodable image", {"word/media/image9.png": b"this is not a png"}, "could not process image")
    office_case("term in a zip entry name", {"word/media/Acme Corp logo.png": png_bytes()}, "still present")

    # a term in a sheet name is not rewritten (it would break formulas) -> must be flagged
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.title = "Acme Corp Data"
    wb.active["A1"] = "Acme Corp"
    xlsx = td / "sheetname.xlsx"
    wb.save(xlsx)
    out = td / "out_xlsx"
    expect_warning("term in an Excel sheet name", run(xlsx, *terms, "--output-dir", out), out, "still present")

    # PDFs: annotations and embedded files are not scrubbed
    import pymupdf
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Acme Corp")
    page.add_text_annot((100, 100), "note mentioning Acme Corp")
    annot_pdf = td / "annot.pdf"
    doc.save(str(annot_pdf))
    out = td / "out_annot"
    expect_warning("PDF annotation", run(annot_pdf, *terms, "--output-dir", out), out, "annotations or form fields")

    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "Acme Corp")
    doc.embfile_add("attachment.txt", b"Acme Corp secret attachment")
    emb_pdf = td / "emb.pdf"
    doc.save(str(emb_pdf))
    out = td / "out_emb"
    expect_warning("PDF embedded file", run(emb_pdf, *terms, "--output-dir", out), out, "embedded files")

    out = td / "out_emb_strip"
    r = run(emb_pdf, *terms, "--strip-metadata", "--output-dir", out)
    result = next(out.iterdir()) if out.exists() and any(out.iterdir()) else None
    count = pymupdf.open(str(result)).embfile_count() if result else -1
    c.check(count == 0 and "embedded files" not in r.stderr, "--strip-metadata removes the embedded file", r.stderr)

    # a clean file must NOT warn (guards against warnings that always fire)
    out = td / "out_clean"
    r = run(base_docx, *terms, "--output-dir", out)
    c.check(r.returncode == 0 and "WARN" not in r.stderr, "clean .docx exits 0 with no warnings", r.stderr)

c.finish()
